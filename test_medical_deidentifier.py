import tempfile
import unittest
import warnings
from pathlib import Path

from medical_deidentifier import MappingStore, main, process_dicom, process_docx, process_xls, redact_text, scan_text, UIDMapper


class DeidentifierTests(unittest.TestCase):
    def test_both_mode_uses_one_patient_id(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            records = root / "records"
            images = root / "images"
            (records / "张三").mkdir(parents=True)
            (images / "张三").mkdir(parents=True)
            (records / "张三" / "病历.txt").write_text("姓名：张三 诊断：肺炎", encoding="utf-8")
            # An unsupported image-side file is enough to verify paired routing; it must go to review.
            (images / "张三" / "note.txt").write_text("not dicom", encoding="utf-8")
            output = root / "output"
            mapping = root / "protected" / "mapping.xlsx"
            rc = main(["--mode", "both", "--records-input", str(records),
                       "--dicom-input", str(images), "--output", str(output),
                       "--mapping", str(mapping)])
            self.assertEqual(rc, 0)
            record_ids = [p.name for p in (output / "records" / "share_ready").iterdir()]
            image_ids = [p.name for p in (output / "dicom" / "share_ready").iterdir()]
            self.assertEqual(record_ids, image_ids)
            self.assertEqual(len(record_ids), 1)

    def test_both_mode_processes_unpaired_folders(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            (root / "records" / "张三").mkdir(parents=True)
            (root / "images" / "李四").mkdir(parents=True)
            (root / "records" / "张三" / "病历.txt").write_text("姓名：张三 诊断：肺炎", encoding="utf-8")
            (root / "images" / "李四" / "note.txt").write_text("not dicom", encoding="utf-8")
            rc = main(["--mode", "both", "--records-input", str(root / "records"),
                       "--dicom-input", str(root / "images"), "--output", str(root / "output"),
                       "--mapping", str(root / "protected" / "mapping.xlsx")])
            self.assertEqual(rc, 0)
            record_ids = [p.name for p in (root / "output" / "records" / "share_ready").iterdir()]
            image_ids = [p.name for p in (root / "output" / "dicom" / "share_ready").iterdir()]
            self.assertEqual(len(record_ids), 1)
            self.assertEqual(len(image_ids), 1)
            self.assertNotEqual(record_ids, image_ids)

    def test_text_redaction_preserves_clinical_numbers(self):
        source = "姓名：张三 诊断：2型糖尿病 血糖7.2 手机13800138000 身份证110101199001011234"
        cleaned, _ = redact_text(source, "张三", "Patient_A1B2C3D4")
        self.assertIn("2型糖尿病", cleaned)
        self.assertIn("血糖7.2", cleaned)
        self.assertFalse(scan_text(cleaned, "张三"))

    def test_mapping_is_stable(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "map.xlsx"
            pid = MappingStore(path).get_or_create("key-one", "张三")
            self.assertEqual(MappingStore(path).get_or_create("key-one", "张三"), pid)

    def test_paired_mapping_reuses_prior_single_mode_id(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "map.xlsx"
            store = MappingStore(path)
            old_pid = store.get_or_create("records-key", "张三")
            paired_pid = store.get_or_create_paired("pair-key", ["records-key", "dicom-key"], "张三")
            self.assertEqual(paired_pid, old_pid)
            self.assertEqual(MappingStore(path).get_or_create("dicom-key", "张三"), old_pid)

    def test_docx_body_table_header_and_metadata(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            src, dst = Path(td) / "a.docx", Path(td) / "b.docx"
            doc = Document()
            doc.add_paragraph("张三 手机13800138000 诊断肺炎")
            doc.add_table(1, 1).cell(0, 0).text = "住院号：ABC12345"
            doc.sections[0].header.paragraphs[0].text = "张三"
            doc.core_properties.author = "张三"
            doc.save(src)
            result = process_docx(src, dst, "张三", "Patient_A1B2C3D4")
            self.assertEqual(result.status, "success")
            check = Document(dst)
            self.assertIn("诊断肺炎", check.paragraphs[0].text)
            self.assertNotIn("张三", check.paragraphs[0].text)
            self.assertEqual(check.core_properties.author, "")

    def test_legacy_xls_cell_redaction_preserves_style(self):
        import xlrd
        import xlwt
        with tempfile.TemporaryDirectory() as td:
            src, dst = Path(td) / "a.xls", Path(td) / "b.xls"
            wb = xlwt.Workbook()
            ws = wb.add_sheet("病历")
            style = xlwt.easyxf("font: bold on; pattern: pattern solid, fore_colour yellow;")
            ws.write(0, 0, "张三 手机13800138000 诊断肺炎", style)
            ws.write(1, 0, 7.2)
            wb.save(str(src))
            result = process_xls(src, dst, "张三", "Patient_A1B2C3D4")
            self.assertEqual(result.status, "success")
            before = xlrd.open_workbook(src, formatting_info=True).sheet_by_index(0)
            after = xlrd.open_workbook(dst, formatting_info=True).sheet_by_index(0)
            self.assertNotIn("张三", after.cell_value(0, 0))
            self.assertIn("诊断肺炎", after.cell_value(0, 0))
            self.assertEqual(after.cell_value(1, 0), 7.2)
            self.assertEqual(before.cell_xf_index(0, 0), after.cell_xf_index(0, 0))

    def test_dicom_mode_redacts_label_and_copies_safe_image(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            patient = root / "input" / "张三"
            patient.mkdir(parents=True)
            label_text = ("病人姓名:  Yan Jin=张三,\nPatient ID:  2025113021288\n"
                          "出生日期:  08 十月 1992\n检查日期:  30 十一月 2025\n"
                          "文件集创建时间:  27 八月 2026\nFileset 创建者：Synapse\n"
                          "文件集内容:  DICOM ONLY\n")
            (patient / "LABEL.TXT").write_text(label_text, encoding="utf-16")
            Image.new("RGB", (2, 2), "white").save(patient / "preview.png")
            rc = main(["--mode", "dicom", "--input", str(root / "input"),
                       "--output", str(root / "output"),
                       "--mapping", str(root / "protected" / "mapping.xlsx")])
            self.assertEqual(rc, 0)
            patient_output = next((root / "output" / "share_ready").iterdir())
            label = (patient_output / "LABEL.TXT").read_text(encoding="utf-16")
            self.assertNotIn("张三", label)
            self.assertNotIn("Yan Jin", label)
            self.assertNotIn("2025113021288", label)
            self.assertNotIn("08 十月 1992", label)
            self.assertNotIn("30 十一月 2025", label)
            self.assertIn("Synapse", label)
            self.assertTrue((patient_output / "preview.png").exists())

    def test_non_target_imaging_assets_are_byte_identical(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            patient = root / "input" / "张三"
            patient.mkdir(parents=True)
            Image.new("RGB", (2, 2), "white").save(patient / "preview.png")
            payload = b"unknown vendor payload\x00\xff"
            (patient / "unknown.bin").write_bytes(payload)
            rc = main(["--mode", "dicom", "--input", str(root / "input"),
                       "--output", str(root / "output"),
                       "--mapping", str(root / "protected" / "mapping.xlsx")])
            self.assertEqual(rc, 0)
            share_patient = next((root / "output" / "share_ready").iterdir())
            self.assertTrue((share_patient / "preview.png").exists())
            self.assertEqual((share_patient / "unknown.bin").read_bytes(), payload)

    def test_dicomdir_reference_tracks_remapped_dicomobj_uid(self):
        import pydicom
        from pydicom.dataset import Dataset, FileDataset, FileMetaDataset
        from pydicom.uid import (ExplicitVRLittleEndian, MediaStorageDirectoryStorage,
                                 SecondaryCaptureImageStorage, generate_uid)
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            patient = root / "input" / "张三"
            obj_dir = patient / "DICOMOBJ"
            obj_dir.mkdir(parents=True)
            sop = generate_uid()

            obj_meta = FileMetaDataset()
            obj_meta.TransferSyntaxUID = ExplicitVRLittleEndian
            obj_meta.MediaStorageSOPClassUID = SecondaryCaptureImageStorage
            obj_meta.MediaStorageSOPInstanceUID = sop
            obj = FileDataset(str(obj_dir / "00000001"), {}, file_meta=obj_meta, preamble=b"\0" * 128)
            obj.SOPClassUID = SecondaryCaptureImageStorage
            obj.SOPInstanceUID = sop
            obj.StudyInstanceUID = generate_uid()
            obj.SeriesInstanceUID = generate_uid()
            obj.PatientName = "张三"
            obj.PatientID = "12345"
            obj.save_as(obj_dir / "00000001", enforce_file_format=True)

            dir_sop = generate_uid()
            dir_meta = FileMetaDataset()
            dir_meta.TransferSyntaxUID = ExplicitVRLittleEndian
            dir_meta.MediaStorageSOPClassUID = MediaStorageDirectoryStorage
            dir_meta.MediaStorageSOPInstanceUID = dir_sop
            dicomdir = FileDataset(str(patient / "DICOMDIR"), {}, file_meta=dir_meta, preamble=b"\0" * 128)
            dicomdir.SOPClassUID = MediaStorageDirectoryStorage
            dicomdir.SOPInstanceUID = dir_sop
            record = Dataset()
            record.DirectoryRecordType = "IMAGE"
            record.PatientName = "张三"
            record.PatientID = "12345"
            record.ReferencedSOPClassUIDInFile = SecondaryCaptureImageStorage
            record.ReferencedSOPInstanceUIDInFile = sop
            record.ReferencedTransferSyntaxUIDInFile = ExplicitVRLittleEndian
            record.ReferencedFileID = ["DICOMOBJ", "00000001"]
            dicomdir.DirectoryRecordSequence = [record]
            dicomdir.save_as(patient / "DICOMDIR", enforce_file_format=True)

            rc = main(["--mode", "dicom", "--input", str(root / "input"),
                       "--output", str(root / "output"),
                       "--mapping", str(root / "protected" / "mapping.xlsx")])
            self.assertEqual(rc, 0)
            out_patient = next((root / "output" / "share_ready").iterdir())
            clean_obj = pydicom.dcmread(out_patient / "DICOMOBJ" / "00000001")
            clean_dir = pydicom.dcmread(out_patient / "DICOMDIR")
            clean_record = clean_dir.DirectoryRecordSequence[0]
            self.assertEqual(clean_record.ReferencedSOPInstanceUIDInFile, clean_obj.SOPInstanceUID)
            self.assertEqual(str(clean_record.PatientID), out_patient.name)
            self.assertEqual(str(clean_record.PatientName), out_patient.name)

    def test_dicom_uid_date_and_pixel_integrity(self):
        import pydicom
        from pydicom.dataset import Dataset, FileDataset, FileMetaDataset
        from pydicom.uid import ExplicitVRLittleEndian, SecondaryCaptureImageStorage, generate_uid
        with tempfile.TemporaryDirectory() as td:
            src, dst = Path(td) / "a.dcm", Path(td) / "b.dcm"
            sop = generate_uid()
            meta = FileMetaDataset()
            meta.TransferSyntaxUID = ExplicitVRLittleEndian
            meta.MediaStorageSOPClassUID = SecondaryCaptureImageStorage
            meta.MediaStorageSOPInstanceUID = sop
            ds = FileDataset(str(src), {}, file_meta=meta, preamble=b"\0" * 128)
            ds.SOPClassUID = SecondaryCaptureImageStorage
            ds.SOPInstanceUID = sop
            ds.StudyInstanceUID = generate_uid()
            ds.SeriesInstanceUID = generate_uid()
            original_failed_uids = [generate_uid(), generate_uid()]
            ds.FailedSOPInstanceUIDList = original_failed_uids
            ds.PatientName = "张三"
            ds.PatientID = "H12345"
            ds.StudyDate = "20200102"
            ds.BurnedInAnnotation = "NO"
            ds.Rows, ds.Columns = 2, 2
            ds.SamplesPerPixel = 1
            ds.PhotometricInterpretation = "MONOCHROME2"
            ds.BitsAllocated = ds.BitsStored = 8
            ds.HighBit = 7
            ds.PixelRepresentation = 0
            ds.PixelData = bytes([0, 1, 2, 3])
            ds.save_as(src, enforce_file_format=True)
            result = process_dicom(src, dst, "张三", "Patient_A1B2C3D4", UIDMapper("secret"), 10)
            self.assertEqual(result.status, "success")
            check = pydicom.dcmread(dst)
            self.assertEqual(str(check.PatientID), "Patient_A1B2C3D4")
            self.assertEqual(check.StudyDate, "20200112")
            self.assertNotEqual(check.SOPInstanceUID, sop)
            self.assertEqual(len(check.FailedSOPInstanceUIDList), 2)
            self.assertNotEqual(list(check.FailedSOPInstanceUIDList), original_failed_uids)
            self.assertEqual(check.PixelData, bytes([0, 1, 2, 3]))

    def test_dicom_empty_pixel_element_does_not_raise_type_error(self):
        import pydicom
        from pydicom.dataset import FileDataset, FileMetaDataset
        from pydicom.uid import ExplicitVRLittleEndian, SecondaryCaptureImageStorage, generate_uid
        with tempfile.TemporaryDirectory() as td:
            src, dst = Path(td) / "empty.dcm", Path(td) / "clean.dcm"
            sop = generate_uid()
            meta = FileMetaDataset()
            meta.TransferSyntaxUID = ExplicitVRLittleEndian
            meta.MediaStorageSOPClassUID = SecondaryCaptureImageStorage
            meta.MediaStorageSOPInstanceUID = sop
            ds = FileDataset(str(src), {}, file_meta=meta, preamble=b"\0" * 128)
            ds.SOPClassUID = SecondaryCaptureImageStorage
            ds.SOPInstanceUID = sop
            ds.StudyInstanceUID = generate_uid()
            ds.SeriesInstanceUID = generate_uid()
            ds.PatientName = "Test"
            ds.PatientID = "123"
            ds.SpecificCharacterSet = ["GB18030", "ISO_IR 100"]
            ds.PixelData = None
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                ds.save_as(src, enforce_file_format=True)
            result = process_dicom(src, dst, "Test", "Patient_A1B2C3D4", UIDMapper("secret"), 10)
            self.assertEqual(result.status, "success")
            clean = pydicom.dcmread(dst)
            self.assertIsNone(getattr(clean, "PixelData", None))
            self.assertEqual(clean.SpecificCharacterSet, "ISO_IR 192")


if __name__ == "__main__":
    unittest.main()
